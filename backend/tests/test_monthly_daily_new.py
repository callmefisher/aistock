"""当月每日新增（monthly_daily_new 步骤）测试 — 无缓存，只用数据日期当日的条件交集数据

覆盖：
- filter_monthly_daily_new_codes 纯函数：
  * 非首交易日：J 含数据日期且当月无更早日期才输出；跨月旧日期不影响输出
  * 首交易日判定（全表 J 列当月只出现数据日期）→ 排除上月任何日期
    （J="2026-10-01,2026-09-10" 排除；J="2026-10-01,2026-08-20" 有效输出）
  * 数据日期为本月非 1 日但当月无其他日期（如长假后首日）→ 同首交易日逻辑
  * 跨年首交易日（1 月 → 排除上一年 12 月）
  * 标绿：J 严格只含唯一日期且等于数据日期；含任何其他日期（含跨月）输出但不标绿
  * 行序去重（首行优先）、J 不含数据日期的行排除
- 数据加载（百日新高 J 列）：DB 优先（空表为权威空）、DB 记录缺 J 列（旧格式）走 Excel 兜底、
  无 DB 记录走 Excel 兜底、主文件查找排除"证券代码"后缀文件
- 数据加载（站上20日线 K 列）：stock_pools DB 优先、Excel sheet"条件交集站上20日线"兜底
- 步骤端到端：输出单列 Excel + 严格新增绿底、幂等重复执行输出一致、
  当日文件被删但 DB 有记录时结果不变（回归一键并行/当日二次执行标记丢失 bug）、
  当日条件交集有数据但无新增 → 成功输出空表、无数据/空表失败、单数字月份目录
- 20日线第二 sheet：单文件双 sheet（"当月每日新增"=百日新高 + "20日线"=站上20日线）、
  同一过滤/标绿逻辑、DB（选股池）优先、主文件无站上20日线 sheet → 跳过该 sheet
  且不影响百日新高主输出

DB 通过 patch core.database.AsyncSessionLocal 注入 FakeSession 隔离，不依赖真实数据库。
"""

import asyncio
import json
import os
import zlib
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest

from services.workflow_executor import WorkflowExecutor, filter_monthly_daily_new_codes


# ---------------------------------------------------------------- DB 假件

def _compress(records) -> bytes:
    return zlib.compress(json.dumps(records, ensure_ascii=False).encode("utf-8"), 6)


class _FakeSession:
    """Async with 协议 + execute 按参数路由：'d'→workflow_results（按日期），
    'n'→stock_pools（按选股池名称）。fetchone 返回单行元组。"""

    def __init__(self, rows, pool_rows=None):
        self._rows = rows            # [(date_str, compressed)] workflow_results
        self._pool_rows = pool_rows or []  # [(name, data_json)] stock_pools

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return False

    async def execute(self, stmt, params=None):
        params = params or {}
        mock_result = MagicMock()
        if "n" in params:  # stock_pools 查询（站上20日线选股池）
            matched = [r for r in self._pool_rows if r[0] == params["n"]]
            mock_result.fetchone = MagicMock(return_value=(matched[0][1],) if matched else None)
        else:  # workflow_results 查询（按日期）
            d = params.get("d", "")
            matched = [r for r in self._rows if r[0] == d]
            mock_result.fetchone = MagicMock(return_value=(matched[0][1],) if matched else None)
        return mock_result


def _fake_session_factory(rows, pool_rows=None):
    def factory(*a, **kw):
        return _FakeSession(rows, pool_rows)
    return factory


def _db_row(date_str, rows_j):
    """构造 workflow_results 记录（内容 = 条件交集主文件首 sheet 行）。
    rows_j: [(证券代码, J列日期串), ...]"""
    j_col = f"2026-03-18至{date_str}期间百日新高的日期"
    records = [
        {"证券代码": c, "证券简称": c + "简称", j_col: j}
        for c, j in rows_j
    ]
    return (date_str, _compress(records))


def _pool_row(date_str, rows_k):
    """构造 stock_pools 记录（"条件交集站上20日线_{date_str}"选股池，data 为 JSON 字符串）。
    rows_k: [(证券代码, K列日期串), ...]"""
    k_col = f"2026-03-18至{date_str}站上20日线日期"
    records = [
        {"证券代码": c, "证券简称": c + "简称", k_col: j}
        for c, j in rows_k
    ]
    return (f"条件交集站上20日线_{date_str}", json.dumps(records, ensure_ascii=False))


@pytest.fixture(autouse=True)
def _fake_db(monkeypatch):
    """默认 DB 无记录 → 文件类测试即 Excel 兜底路径，同时避免真实 DB 连接。"""
    monkeypatch.setattr("core.database.AsyncSessionLocal", _fake_session_factory([]))


# ---------------------------------------------------------------- 公共

def _make_intersection_file(base_dir, date_str, rows, ma20_rows=None):
    """生成条件交集主文件（sheet"条件交集百日新高"），rows: [(证券代码, J列日期串), ...]
    ma20_rows 非 None 时追加 sheet"条件交集站上20日线"（K列日期），ma20_rows: [(证券代码, K列日期串), ...]
    同时放置"证券代码"后缀干扰文件，验证主文件查找的排除逻辑"""
    day_dir = os.path.join(base_dir, "条件交集", date_str)
    os.makedirs(day_dir, exist_ok=True)
    compact = date_str.replace("-", "")
    j_col = f"2026-03-18至{date_str}期间百日新高的日期"
    df = pd.DataFrame({
        "证券代码": [r[0] for r in rows],
        "证券简称": [r[0] + "简称" for r in rows],
        j_col: [r[1] for r in rows],
    })
    main_file = os.path.join(day_dir, f"7条件交集{compact}.xlsx")
    with pd.ExcelWriter(main_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="条件交集百日新高", index=False)
        if ma20_rows is not None:
            k_col = f"2026-03-18至{date_str}站上20日线日期"
            df_m = pd.DataFrame({
                "证券代码": [r[0] for r in ma20_rows],
                "证券简称": [r[0] + "简称" for r in ma20_rows],
                f"2026-03-18至{date_str}站上20日线次数": ["1"] * len(ma20_rows),
                k_col: [r[1] for r in ma20_rows],
            })
            df_m.to_excel(writer, sheet_name="条件交集站上20日线", index=False)
    pd.DataFrame({"证券代码": []}).to_excel(
        os.path.join(day_dir, f"7条件交集{compact}百日新高证券代码.xlsx"),
        index=False, engine="openpyxl",
    )
    return main_file


def _run_step(tmp_path, config):
    executor = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    return asyncio.run(executor.execute_step("monthly_daily_new", config))


def _read_output(path, sheet_name=None):
    """读取输出Excel指定sheet（默认首个sheet）→ (表头, 值列表, 填充色列表)"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name else wb.active
    values = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    fills = []
    for r in range(2, ws.max_row + 1):
        f = ws.cell(r, 1).fill
        fills.append(f.start_color.rgb if f.fill_type == "solid" else None)
    return ws.cell(1, 1).value, values, fills


# ---------------------------------------------------------------- 纯函数过滤

def test_filter_normal_day():
    """非首交易日：J 含数据日期且当月无更早日期才输出；跨月旧日期不影响输出，但影响标绿"""
    rows = [
        ("000001.SZ", "2026-10-01,2026-10-09"),   # 当月更早日期 → 排除
        ("000002.SZ", "2026-10-09"),               # 纯新 → 输出+标绿
        ("000003.SZ", "2026-10-09,2026-09-28"),    # 当月新增+跨月旧日期 → 输出不标绿
        ("000004.SZ", "2026-10-03"),               # J 不含数据日期 → 排除
        ("000002.SZ", "2026-10-09"),               # 重复代码 → 去重
    ]
    out, marked, first = filter_monthly_daily_new_codes(rows, "2026-10-09")
    assert out == ["000002.SZ", "000003.SZ"]
    assert marked == ["000002.SZ"]
    assert first is False


def test_filter_first_trading_day_on_month_1st():
    """首交易日（数据日期为本月1日）：排除上月日期；上上月日期有效"""
    rows = [
        ("000001.SZ", "2026-10-01"),               # 纯新 → 输出+标绿
        ("000002.SZ", "2026-10-01,2026-09-10"),     # 上月(9月)日期 → 排除
        ("000003.SZ", "2026-10-01,2026-08-20"),     # 上上月(8月)日期 → 有效输出，不标绿
    ]
    out, marked, first = filter_monthly_daily_new_codes(rows, "2026-10-01")
    assert first is True
    assert out == ["000001.SZ", "000003.SZ"]
    assert marked == ["000001.SZ"]


def test_filter_first_trading_day_later_in_month():
    """数据日期非1日但全表当月只出现数据日期（如长假后首日）→ 按首交易日逻辑"""
    rows = [
        ("000001.SZ", "2026-10-09"),
        ("000002.SZ", "2026-10-09,2026-09-30"),   # 上月(9月)日期 → 排除
        ("000003.SZ", "2026-10-09,2026-08-15"),   # 8月日期 → 有效输出
    ]
    out, marked, first = filter_monthly_daily_new_codes(rows, "2026-10-09")
    assert first is True
    assert out == ["000001.SZ", "000003.SZ"]
    assert marked == ["000001.SZ"]


def test_filter_first_trading_day_year_boundary():
    """跨年首交易日：1月 → 排除上一年 12 月，11 月及更早有效"""
    rows = [
        ("000001.SZ", "2027-01-04"),
        ("000002.SZ", "2027-01-04,2026-12-31"),   # 上月(2026-12) → 排除
        ("000003.SZ", "2027-01-04,2026-11-20"),   # 2026-11 → 有效输出
    ]
    out, marked, first = filter_monthly_daily_new_codes(rows, "2027-01-04")
    assert first is True
    assert out == ["000001.SZ", "000003.SZ"]
    assert marked == ["000001.SZ"]


def test_filter_empty_month_dates():
    """全表 J 列均无当月日期 → 首交易日判定为真，但无任何输出（均不含数据日期）"""
    rows = [("000001.SZ", "2026-09-20")]
    out, marked, first = filter_monthly_daily_new_codes(rows, "2026-10-09")
    assert out == [] and marked == [] and first is True


# ---------------------------------------------------------------- 数据加载

def test_load_rows_db_priority(tmp_path):
    """DB 有记录（无 Excel 文件）→ 直接用 DB，不需 Excel"""
    rows = [_db_row("2026-10-09", [
        ("000001.SZ", "2026-10-09"),
        ("000002.SZ", "2026-10-01,2026-10-09"),
    ])]
    ex = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(rows)):
        loaded = asyncio.run(ex._load_intersection_rows(
            os.path.join(str(tmp_path), "条件交集"), "2026-10-09"))
    assert loaded == [("000001.SZ", "2026-10-09"), ("000002.SZ", "2026-10-01,2026-10-09")]


def test_load_rows_db_empty_authoritative(tmp_path):
    """DB 记录为空列表 → 权威空表，不走 Excel 兜底"""
    rows = [("2026-10-09", _compress([]))]
    _make_intersection_file(str(tmp_path), "2026-10-09", [("000005.SZ", "2026-10-09")])
    ex = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(rows)):
        loaded = asyncio.run(ex._load_intersection_rows(
            os.path.join(str(tmp_path), "条件交集"), "2026-10-09"))
    assert loaded == []


def test_load_rows_db_old_format_fallback_excel(tmp_path):
    """DB 记录缺 J 列字段（旧格式首 sheet）→ 走 Excel 兜底"""
    rows = [("2026-10-09", _compress([{"证券代码": "000001.SZ", "其他": "x"}]))]
    _make_intersection_file(str(tmp_path), "2026-10-09", [("000005.SZ", "2026-10-09")])
    ex = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(rows)):
        loaded = asyncio.run(ex._load_intersection_rows(
            os.path.join(str(tmp_path), "条件交集"), "2026-10-09"))
    assert loaded == [("000005.SZ", "2026-10-09")]


def test_load_rows_nothing(tmp_path):
    """DB 与 Excel 均无数据 → None"""
    ex = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    loaded = asyncio.run(ex._load_intersection_rows(
        os.path.join(str(tmp_path), "条件交集"), "2026-10-09"))
    assert loaded is None


# ---------------------------------------------------------------- 站上20日线数据加载（K列）

def test_load_rows_ma20_db_priority(tmp_path):
    """stock_pools 有选股池记录（无 Excel 文件）→ 直接用 DB"""
    pools = [_pool_row("2026-10-09", [
        ("000001.SZ", "2026-10-09"),
        ("000002.SZ", "2026-10-01,2026-10-09"),
    ])]
    ex = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    with patch("core.database.AsyncSessionLocal", _fake_session_factory([], pools)):
        loaded = asyncio.run(ex._load_intersection_rows_ma20(
            os.path.join(str(tmp_path), "条件交集"), "2026-10-09"))
    assert loaded == [("000001.SZ", "2026-10-09"), ("000002.SZ", "2026-10-01,2026-10-09")]


def test_load_rows_ma20_excel_fallback(tmp_path):
    """DB 无选股池记录 → Excel 主文件 sheet"条件交集站上20日线"K列兜底"""
    _make_intersection_file(str(tmp_path), "2026-10-09",
                            [("000001.SZ", "2026-10-09")],
                            ma20_rows=[("000005.SZ", "2026-10-09")])
    ex = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    loaded = asyncio.run(ex._load_intersection_rows_ma20(
        os.path.join(str(tmp_path), "条件交集"), "2026-10-09"))
    assert loaded == [("000005.SZ", "2026-10-09")]


def test_load_rows_ma20_no_sheet(tmp_path):
    """主文件无"条件交集站上20日线"sheet（旧格式）→ None"""
    _make_intersection_file(str(tmp_path), "2026-10-09", [("000001.SZ", "2026-10-09")])
    ex = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    loaded = asyncio.run(ex._load_intersection_rows_ma20(
        os.path.join(str(tmp_path), "条件交集"), "2026-10-09"))
    assert loaded is None


# ---------------------------------------------------------------- monthly_daily_new 步骤

def test_monthly_daily_new_step(tmp_path):
    """Excel 兜底：输出仅当日新增代码，J 严格单一日期标绿；幂等重复执行输出一致"""
    base = str(tmp_path)
    _make_intersection_file(base, "2026-10-09", [
        ("000001.SZ", "2026-10-01,2026-10-09"),   # 当月更早 → 排除
        ("000002.SZ", "2026-10-09"),               # 纯新 → 输出+标绿
        ("000003.SZ", "2026-10-09,2026-09-28"),    # 当月新增+跨月旧 → 输出不标绿
        ("000004.SZ", "2026-10-03,2026-10-09"),   # 当月更早 → 排除
    ])

    result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is True
    assert os.path.basename(result["file_path"]) == "17 10月新增1009.xlsx"
    # 数据目录按月组织
    assert os.path.basename(os.path.dirname(result["file_path"])) == "2026年10月新增"
    assert result["new_codes"] == ["000002.SZ", "000003.SZ"]
    assert result["marked_codes"] == ["000002.SZ"]

    header, values, fills = _read_output(result["file_path"])
    assert header == "证券代码"
    assert values == ["000002.SZ", "000003.SZ"]
    assert fills == ["FFD8EAD8", None]   # 严格单一日期标绿，跨月旧日期不标绿

    # 幂等：重复执行输出一致
    result2 = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result2["new_codes"] == result["new_codes"]
    assert result2["marked_codes"] == result["marked_codes"]
    _, values2, fills2 = _read_output(result2["file_path"])
    assert values2 == values and fills2 == fills


def test_monthly_daily_new_first_trading_day_step(tmp_path):
    """首交易日（本月1日）：排除上月日期，上上月日期有效；消息提示首交易日"""
    base = str(tmp_path)
    _make_intersection_file(base, "2026-10-01", [
        ("000001.SZ", "2026-10-01"),
        ("000002.SZ", "2026-10-01,2026-09-10"),   # 上月日期 → 排除
        ("000003.SZ", "2026-10-01,2026-08-20"),   # 上上月 → 输出不标绿
    ])

    result = _run_step(tmp_path, {"date_str": "2026-10-01"})
    assert result["success"] is True
    assert os.path.basename(result["file_path"]) == "17 10月新增1001.xlsx"
    assert result["new_codes"] == ["000001.SZ", "000003.SZ"]
    assert result["marked_codes"] == ["000001.SZ"]
    assert "当月第一个交易日" in result["message"]

    _, values, fills = _read_output(result["file_path"])
    assert values == ["000001.SZ", "000003.SZ"]
    assert fills == ["FFD8EAD8", None]


def test_monthly_daily_new_today_from_db(tmp_path):
    """回归 bug：条件交集再次执行时先删除了当日文件且失败，文件丢失；
    但首次成功的结果已入库 → 当月每日新增从 DB 读到当日数据，输出/标记正常"""
    rows = [_db_row("2026-10-09", [
        ("000001.SZ", "2026-10-01,2026-10-09"),
        ("000002.SZ", "2026-10-09"),
        ("000003.SZ", "2026-10-09"),
    ])]
    # 当日 Excel 文件不存在（已被删除）
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(rows)):
        result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is True
    assert result["new_codes"] == ["000002.SZ", "000003.SZ"]
    assert result["marked_codes"] == ["000002.SZ", "000003.SZ"]

    _, values, fills = _read_output(result["file_path"])
    assert values == ["000002.SZ", "000003.SZ"]
    assert fills == ["FFD8EAD8", "FFD8EAD8"]


def test_monthly_daily_new_no_new_codes(tmp_path):
    """当日条件交集有数据但无新增 → 成功输出空表（仅表头）"""
    base = str(tmp_path)
    _make_intersection_file(base, "2026-10-09", [
        ("000001.SZ", "2026-10-03,2026-10-09"),   # 当月更早 → 无新增
    ])
    result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is True
    assert result["new_codes"] == []
    assert "无新增" in result["message"]
    header, values, fills = _read_output(result["file_path"])
    assert header == "证券代码" and values == [] and fills == []


def test_monthly_daily_new_empty_table(tmp_path):
    """当日条件交集为空表（DB 空记录）→ 失败并提示"""
    rows = [("2026-10-09", _compress([]))]
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(rows)):
        result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is False
    assert "空表" in result["message"]


def test_monthly_daily_new_no_data(tmp_path):
    """当月无任何数据（DB/Excel 均无）→ 失败并提示"""
    result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is False
    assert "无条件交集数据" in result["message"]


def test_monthly_daily_new_single_digit_month(tmp_path):
    """单数字月份：目录 2026年3月新增 / 文件 17 3月新增0305.xlsx"""
    _make_intersection_file(str(tmp_path), "2026-03-05", [("600000.SH", "2026-03-05")])
    result = _run_step(tmp_path, {"date_str": "2026-03-05"})
    assert result["success"] is True
    rel = os.path.relpath(result["file_path"], str(tmp_path)).replace(os.sep, "/")
    assert rel == "当月每日新增/2026年3月新增/17 3月新增0305.xlsx"


# ---------------------------------------------------------------- 20日线第二sheet

def test_monthly_daily_new_ma20_output(tmp_path):
    """主文件含两个 sheet → 单文件双 sheet："当月每日新增"=百日新高新增 +
    "20日线"=站上20日线新增，两路过滤/标绿逻辑一致"""
    base = str(tmp_path)
    _make_intersection_file(base, "2026-10-09", [
        ("000001.SZ", "2026-10-01,2026-10-09"),   # 当月更早 → 排除
        ("000002.SZ", "2026-10-09"),               # 纯新 → 输出+标绿
    ], ma20_rows=[
        ("000001.SZ", "2026-10-09"),               # 纯新 → 输出+标绿
        ("000003.SZ", "2026-10-09,2026-10-03"),    # 当月更早 → 排除
        ("000004.SZ", "2026-10-09,2026-09-28"),    # 当月新增+跨月旧 → 输出不标绿
    ])

    result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is True
    out_file = result["file_path"]
    assert os.path.basename(out_file) == "17 10月新增1009.xlsx"   # 仍只有一个输出文件
    assert result["new_codes"] == ["000002.SZ"]
    assert "20日线新增2个代码（标绿1个）" in result["message"]

    # sheet 结构：主sheet在前，20日线sheet在后
    from openpyxl import load_workbook
    wb = load_workbook(out_file)
    assert wb.sheetnames == ["当月每日新增", "20日线"]

    # 主 sheet（百日新高）
    header, values, fills = _read_output(out_file)
    assert header == "证券代码"
    assert values == ["000002.SZ"]
    assert fills == ["FFD8EAD8"]

    # 20日线 sheet（同一套过滤/标绿逻辑）
    header_m, values_m, fills_m = _read_output(out_file, "20日线")
    assert header_m == "证券代码"
    assert values_m == ["000001.SZ", "000004.SZ"]
    assert fills_m == ["FFD8EAD8", None]

    # 幂等：重复执行输出一致
    result2 = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result2["new_codes"] == result["new_codes"]
    wb2 = load_workbook(result2["file_path"])
    assert wb2.sheetnames == ["当月每日新增", "20日线"]
    _, values2, fills2 = _read_output(result2["file_path"], "20日线")
    assert values2 == values_m and fills2 == fills_m


def test_monthly_daily_new_ma20_from_db(tmp_path):
    """20日线数据源 DB（stock_pools 选股池）优先：Excel 无站上20日线 sheet 时仍写入20日线 sheet"""
    pools = [_pool_row("2026-10-09", [
        ("000001.SZ", "2026-10-09"),
        ("000002.SZ", "2026-10-01,2026-10-09"),   # 当月更早 → 排除
    ])]
    # 主文件只有百日新高 sheet（无站上20日线）
    _make_intersection_file(str(tmp_path), "2026-10-09", [("000002.SZ", "2026-10-09")])
    with patch("core.database.AsyncSessionLocal", _fake_session_factory([], pools)):
        result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is True
    from openpyxl import load_workbook
    assert load_workbook(result["file_path"]).sheetnames == ["当月每日新增", "20日线"]
    _, values, fills = _read_output(result["file_path"], "20日线")
    assert values == ["000001.SZ"]
    assert fills == ["FFD8EAD8"]


def test_monthly_daily_new_ma20_missing_skipped(tmp_path):
    """主文件无站上20日线 sheet 且 DB 无选股池 → 仅主 sheet，百日新高输出不受影响"""
    base = str(tmp_path)
    _make_intersection_file(base, "2026-10-09", [("000002.SZ", "2026-10-09")])
    result = _run_step(tmp_path, {"date_str": "2026-10-09"})
    assert result["success"] is True
    assert result["new_codes"] == ["000002.SZ"]
    assert "20日线数据缺失" in result["message"]

    from openpyxl import load_workbook
    assert load_workbook(result["file_path"]).sheetnames == ["当月每日新增"]
    # 月度目录下只有这一个输出文件
    out_dir = os.path.join(base, "当月每日新增", "2026年10月新增")
    assert sorted(os.listdir(out_dir)) == ["17 10月新增1009.xlsx"]


def test_resolver_month_dir(tmp_path):
    """下载端点依赖 _get_daily_dir：新类型解析为月度目录，其他类型不受影响"""
    executor = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="当月每日新增")
    assert executor._get_daily_dir("2026-10-09") == os.path.join(str(tmp_path), "当月每日新增", "2026年10月新增")
    assert executor._get_daily_dir("2026-03-05") == os.path.join(str(tmp_path), "当月每日新增", "2026年3月新增")
    ex2 = WorkflowExecutor(base_dir=str(tmp_path), workflow_type="条件交集")
    assert ex2._get_daily_dir("2026-10-09") == os.path.join(str(tmp_path), "条件交集", "2026-10-09")

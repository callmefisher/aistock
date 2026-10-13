"""本月AI筛选20日线汇总（monthly_ai_ma20_summary 步骤）+ 双sheet入库 测试

覆盖：
- save_workflow_result「当月每日新增」双 sheet 入库：
  * 首 sheet → step_type='final'；第 2 个 sheet「20日线」→ step_type='final_ma20'
  * 文件无「20日线」sheet → 只存 final，不报错
  * 非「当月每日新增」类型 → 行为不变（只存 final）
- execute_monthly_ai_ma20_summary 端到端（DB 用 FakeSession 注入，Excel 兜底用真实文件）：
  * DB 优先读「当月每日新增」final_ma20（SQL 按 step_type='final_ma20' 过滤）；
    无 DB 记录走 Excel 第 2 个 sheet「20日线」兜底（不读首 sheet）；
    Excel 无「20日线」sheet / DB 与 Excel 均无 → 失败提示
  * 输出到 data/excel/本月AI筛选20日线汇总/{YYYY-MM}/19 {M}月AI筛选20日线汇总{MMDD}.xlsx
  * 日期 sheet（MMDD 命名）强制覆盖；Sheet1 双表双图增量更新（复用 monthly_ai_summary 逻辑）
  * 首次新建；再次执行（新数据日期，无同名）→ 复制最新文件为新名；重复执行当日 → 复用同名仅刷新当日列

DB 通过 patch core.database.AsyncSessionLocal（延迟导入）/ services.workflow_result_service.AsyncSessionLocal
注入 FakeSession 隔离，不依赖真实数据库。
"""

import json
import os
import zlib
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from services.monthly_ai_ma20_summary import execute_monthly_ai_ma20_summary
from services.monthly_ai_summary import _all_a_sector_cache
from services.workflow_executor import WorkflowExecutor
from services.workflow_result_service import save_workflow_result


# ---------------------------------------------------------------- 工具

def _compress(records) -> bytes:
    return zlib.compress(json.dumps(records, ensure_ascii=False).encode("utf-8"), 6)


class _FakeSession:
    """Async with 协议 + execute 按 'd' 参数返回 workflow_results 单行；记录 SQL 供断言"""

    def __init__(self, rows):
        self._rows = rows  # [(date_str, compressed)]
        self.executed_sql = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return False

    async def execute(self, stmt, params=None):
        self.executed_sql.append(str(stmt))
        params = params or {}
        mock_result = MagicMock()
        d = params.get("d", "")
        matched = [r for r in self._rows if r[0] == d]
        mock_result.fetchone = MagicMock(return_value=(matched[0][1],) if matched else None)
        return mock_result


def _fake_session_factory(rows):
    def factory(*a, **kw):
        return _FakeSession(rows)
    return factory


class _UpsertCaptureSession:
    """捕获 save_workflow_result 的 UPSERT 参数（step_type/records）"""

    def __init__(self, calls):
        self._calls = calls  # 共享 list，跨多次 async with 累积

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return False

    async def execute(self, stmt, params=None):
        self._calls.append(dict(params or {}))
        return MagicMock()

    async def commit(self):
        pass


def _capture_factory(calls):
    def factory(*a, **kw):
        return _UpsertCaptureSession(calls)
    return factory


def _write_all_a_file(base_dir, codes_sectors):
    """写全部A股文件：多行表头（含换行单元格）1级/2级行业列"""
    os.makedirs(os.path.join(base_dir, "全部A股"), exist_ok=True)
    l1_header = "所属行业名称\n[行业标准]申万行业(2021)\n[行业类别]1级\n[截止日期]最新"
    l2_header = "所属行业名称\n[行业标准]申万行业(2021)\n[行业类别]2级\n[截止日期]最新"
    df = pd.DataFrame({
        "证券代码": [c for c, _, _ in codes_sectors],
        "证券简称": [f"股{i}" for i in range(len(codes_sectors))],
        l1_header: [l1 for _, l1, _ in codes_sectors],
        l2_header: [l2 for _, _, l2 in codes_sectors],
    })
    df.to_excel(os.path.join(base_dir, "全部A股", "6、全部A股 板块信息2026 0901.xlsx"), index=False)


def _write_daily_new_excel(base_dir, date_str, codes, ma20_codes=None):
    """写「当月每日新增」输出文件：单/双 sheet（首sheet=当月每日新增，第2个sheet=20日线）"""
    d = os.path.join(base_dir, "当月每日新增", f"{date_str[:4]}年{int(date_str[5:7])}月新增")
    os.makedirs(d, exist_ok=True)
    mmdd = date_str[5:7] + date_str[8:10]
    path = os.path.join(d, f"17 {int(date_str[5:7])}月新增{mmdd}.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame({"证券代码": codes}).to_excel(w, index=False, sheet_name="当月每日新增")
        if ma20_codes is not None:
            pd.DataFrame({"证券代码": ma20_codes}).to_excel(w, index=False, sheet_name="20日线")
    return path


@pytest.fixture(autouse=True)
def _clear_caches():
    _all_a_sector_cache.clear()
    yield
    _all_a_sector_cache.clear()


# ---------------------------------------------------------------- 双 sheet 入库（save_workflow_result）

@pytest.mark.asyncio
async def test_save_result_daily_new_dual_sheets(tmp_path):
    """「当月每日新增」双 sheet：final=首sheet + final_ma20=20日线sheet，两条记录"""
    path = _write_daily_new_excel(str(tmp_path), "2026-10-12",
                                  ["000009.SZ"], ma20_codes=["600000.SH", "000001.SZ"])
    calls = []
    with patch("services.workflow_result_service.AsyncSessionLocal", _capture_factory(calls)):
        ok = await save_workflow_result(1, "当月每日新增", "wf", "2026-10-12", path)
    assert ok
    assert [c["step_type"] for c in calls] == ["final", "final_ma20"]

    def _records(c):
        return json.loads(zlib.decompress(c["data_compressed"]).decode("utf-8"))

    assert [r["证券代码"] for r in _records(calls[0])] == ["000009.SZ"]
    assert [r["证券代码"] for r in _records(calls[1])] == ["600000.SH", "000001.SZ"]
    assert calls[0]["date_str"] == calls[1]["date_str"] == "2026-10-12"


@pytest.mark.asyncio
async def test_save_result_daily_new_without_ma20_sheet(tmp_path):
    """文件无「20日线」sheet（当日20日线数据缺失）→ 只存 final，不报错"""
    path = _write_daily_new_excel(str(tmp_path), "2026-10-12", ["000009.SZ"], ma20_codes=None)
    calls = []
    with patch("services.workflow_result_service.AsyncSessionLocal", _capture_factory(calls)):
        ok = await save_workflow_result(1, "当月每日新增", "wf", "2026-10-12", path)
    assert ok
    assert [c["step_type"] for c in calls] == ["final"]


@pytest.mark.asyncio
async def test_save_result_other_type_single_sheet(tmp_path):
    """非「当月每日新增」类型 → 行为不变（只存 final）"""
    path = os.path.join(tmp_path, "1并购重组20261012.xlsx")
    pd.DataFrame({"证券代码": ["000009.SZ"]}).to_excel(path, index=False)
    calls = []
    with patch("services.workflow_result_service.AsyncSessionLocal", _capture_factory(calls)):
        ok = await save_workflow_result(1, "并购重组", "wf", "2026-10-12", path)
    assert ok
    assert [c["step_type"] for c in calls] == ["final"]


# ---------------------------------------------------------------- 主步骤端到端

def _executor(base_dir):
    return WorkflowExecutor(base_dir=base_dir, workflow_type="本月AI筛选20日线汇总")


async def _run(executor, date_str, output_filename=""):
    return await execute_monthly_ai_ma20_summary(
        executor, {"date_str": date_str, "output_filename": output_filename}, date_str
    )


def _out_path(base, date_str):
    mmdd = date_str[5:7] + date_str[8:10]
    return os.path.join(base, "本月AI筛选20日线汇总", date_str[:7],
                        f"19 {int(date_str[5:7])}月AI筛选20日线汇总{mmdd}.xlsx")


@pytest.mark.asyncio
async def test_ma20_summary_db_first_and_new_file(tmp_path):
    """DB 优先（final_ma20）：输出目录/文件名正确，日期sheet + Sheet1双表双图"""
    base = str(tmp_path)
    _write_all_a_file(base, [
        ("000009.SZ", "机械设备", "电池"),
        ("000001.SZ", "医药生物", "化学制药"),
        ("600000.SH", "医药生物", "国有大型银行"),
    ])
    db_rows = [("2026-10-12", _compress([{"证券代码": "000009.SZ"}, {"证券代码": "000001.SZ"}, {"证券代码": "600000.SH"}]))]

    with patch("core.database.AsyncSessionLocal", _fake_session_factory(db_rows)):
        result = await _run(_executor(base), "2026-10-12")
    assert result["success"], result.get("message")

    out = _out_path(base, "2026-10-12")
    assert os.path.exists(out), f"输出文件不存在: {out}"

    wb = load_workbook(out)
    assert wb.sheetnames == ["Sheet1", "1012"], wb.sheetnames

    # 日期sheet内容：20日线代码 + 全部A股板块匹配
    ws = wb["1012"]
    assert [c.value for c in ws[1]] == ["证券代码", "一级板块", "二级板块"]
    rows = {ws.cell(r, 1).value: (ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(2, ws.max_row + 1)}
    assert rows["000009.SZ"] == ("机械设备", "电池")
    assert rows["600000.SH"] == ("医药生物", "国有大型银行")

    # Sheet1 统计表 + 两张折线图
    ws1 = wb["Sheet1"]
    assert ws1.cell(1, 1).value == "一级板块"
    assert ws1.cell(1, 2).value == "10月12日"
    stats = {ws1.cell(r, 1).value: ws1.cell(r, 2).value for r in range(2, ws1.max_row + 1) if ws1.cell(r, 1).value}
    assert stats["医药生物"] == 2 and stats["机械设备"] == 1
    assert len(ws1._charts) == 2

    # DB 查询按 step_type='final_ma20' 过滤（区别于首 sheet 的 final）
    import inspect
    from services import monthly_ai_ma20_summary as m
    assert "final_ma20" in inspect.getsource(m._load_ma20_records_from_db)


@pytest.mark.asyncio
async def test_ma20_summary_excel_fallback_reads_second_sheet(tmp_path):
    """无 DB 记录 → Excel 兜底：只读第2个sheet「20日线」（首sheet代码不参与）"""
    base = str(tmp_path)
    _write_all_a_file(base, [("000009.SZ", "机械设备", "电池")])
    # 首 sheet 与 20日线 sheet 代码不同，验证读的是 20日线
    _write_daily_new_excel(base, "2026-10-12", ["600000.SH"], ma20_codes=["000009.SZ"])

    with patch("core.database.AsyncSessionLocal", _fake_session_factory([])):
        result = await _run(_executor(base), "2026-10-12")
    assert result["success"], result.get("message")

    ws = load_workbook(_out_path(base, "2026-10-12"))["1012"]
    codes = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert codes == ["000009.SZ"]  # 20日线 sheet 的代码，而非首 sheet 的 600000.SH
    assert ws.cell(2, 2).value == "机械设备"


@pytest.mark.asyncio
async def test_ma20_summary_missing_ma20_sheet_and_no_data(tmp_path):
    """Excel 文件存在但无「20日线」sheet → 失败；DB与Excel均无 → 失败且不产出文件"""
    base = str(tmp_path)
    _write_all_a_file(base, [("000009.SZ", "机械设备", "电池")])
    _write_daily_new_excel(base, "2026-10-12", ["000009.SZ"], ma20_codes=None)  # 无20日线sheet

    with patch("core.database.AsyncSessionLocal", _fake_session_factory([])):
        result = await _run(_executor(base), "2026-10-12")
    assert not result["success"]
    assert "20日线" in result["message"]

    with patch("core.database.AsyncSessionLocal", _fake_session_factory([])):
        result = await _run(_executor(base), "2026-10-13")
    assert not result["success"]
    assert "当月每日新增" in result["message"]
    assert not os.path.exists(_out_path(base, "2026-10-13"))


@pytest.mark.asyncio
async def test_ma20_summary_empty_day(tmp_path):
    """20日线当日空表（DB空记录[]）→ 合法，输出仅表头的日期sheet"""
    base = str(tmp_path)
    db_rows = [("2026-10-12", _compress([]))]
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(db_rows)):
        result = await _run(_executor(base), "2026-10-12")
    assert result["success"], result.get("message")

    wb = load_workbook(_out_path(base, "2026-10-12"))
    assert wb["1012"].max_row == 1  # 仅表头
    assert wb["Sheet1"].cell(1, 1).value == "一级板块"


@pytest.mark.asyncio
async def test_ma20_summary_copy_latest_and_incremental(tmp_path):
    """1008先建（新建）；1012执行（无同名）→ 复制1008为1012：Sheet1增量、sheet降序、
    1008原样保留；重复执行1012 → 复用同名仅刷新当日列"""
    base = str(tmp_path)
    _write_all_a_file(base, [
        ("000009.SZ", "机械设备", "电池"),
        ("000001.SZ", "医药生物", "化学制药"),
    ])
    d1008 = [("2026-10-08", _compress([{"证券代码": "000009.SZ"}, {"证券代码": "000001.SZ"}]))]
    d1012 = [("2026-10-12", _compress([{"证券代码": "000001.SZ"}]))]
    ex = _executor(base)

    with patch("core.database.AsyncSessionLocal", _fake_session_factory(d1008)):
        r1 = await _run(ex, "2026-10-08")
    assert r1["success"] and "新建" in r1["message"]
    assert os.path.exists(_out_path(base, "2026-10-08"))

    with patch("core.database.AsyncSessionLocal", _fake_session_factory(d1012)):
        r2 = await _run(ex, "2026-10-12")
    assert r2["success"], r2.get("message")
    assert "复制" in r2["message"]

    # 目录两个文件：1008保留 + 1012新副本
    month_dir = os.path.join(base, "本月AI筛选20日线汇总", "2026-10")
    files = sorted(os.listdir(month_dir))
    assert files == ["19 10月AI筛选20日线汇总1008.xlsx", "19 10月AI筛选20日线汇总1012.xlsx"], files

    # 1012 副本：Sheet1 增量（10月8日列沿用 + 10月12日新列），sheet 降序
    wb = load_workbook(_out_path(base, "2026-10-12"))
    assert wb.sheetnames == ["Sheet1", "1012", "1008"], wb.sheetnames
    ws1 = wb["Sheet1"]
    assert [ws1.cell(1, c).value for c in (1, 2, 3)] == ["一级板块", "10月8日", "10月12日"]
    stats = {ws1.cell(r, 1).value: (ws1.cell(r, 2).value, ws1.cell(r, 3).value)
             for r in range(2, ws1.max_row + 1) if ws1.cell(r, 1).value}
    assert stats["医药生物"] == (1, 2)   # 1008累计1 + 1012新增1 = 2
    assert stats["机械设备"] == (1, 1)   # 1008累计1 + 1012新增0 = 1

    # 1008 原文件保持原样
    wb_1008 = load_workbook(_out_path(base, "2026-10-08"))
    assert wb_1008.sheetnames == ["Sheet1", "1008"], wb_1008.sheetnames

    # 重复执行1012（数据变为2个代码）→ 复用同名，仅刷新10月12日列
    d1012b = [("2026-10-12", _compress([{"证券代码": "000009.SZ"}, {"证券代码": "000001.SZ"}]))]
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(d1012b)):
        r3 = await _run(ex, "2026-10-12")
    assert r3["success"] and "复用" in r3["message"]
    wb = load_workbook(_out_path(base, "2026-10-12"))
    ws1 = wb["Sheet1"]
    stats = {ws1.cell(r, 1).value: (ws1.cell(r, 2).value, ws1.cell(r, 3).value)
             for r in range(2, ws1.max_row + 1) if ws1.cell(r, 1).value}
    assert stats["医药生物"] == (1, 2)   # 当日列刷新 = 前列累计1 + 当日新增1 = 2
    assert stats["机械设备"] == (1, 2)   # 当日列刷新 = 前列累计1 + 当日新增1 = 2
    # 目录未新增文件
    assert sorted(os.listdir(month_dir)) == files


@pytest.mark.asyncio
async def test_ma20_summary_copy_latest_by_filename_date(tmp_path):
    """目录非空无同名 → 复制「文件名日期最新」的文件（mtime 干扰不影响）"""
    from openpyxl import Workbook

    base = str(tmp_path)
    _write_all_a_file(base, [("000009.SZ", "机械设备", "电池")])
    month_dir = os.path.join(base, "本月AI筛选20日线汇总", "2027-03")
    os.makedirs(month_dir, exist_ok=True)
    row = ("000009.SZ", "机械设备", "电池")
    # 预置逐日累积文件（无 Sheet1，模拟手工/旧版产物）
    for dd, sheets in [
        ("0305", [("0305", [row])]),
        ("0306", [("0306", [row]), ("0305", [row])]),
        ("0307", [("0307", [row]), ("0306", [row]), ("0305", [row])]),
    ]:
        wb = Workbook()
        wb.remove(wb.active)
        for sn, rows_ in sheets:
            ws = wb.create_sheet(sn)
            ws.append(["证券代码", "一级板块", "二级板块"])
            for r_ in rows_:
                ws.append(list(r_))
        wb.save(os.path.join(month_dir, f"19 3月AI筛选20日线汇总{dd}.xlsx"))
    # mtime 干扰：0305 修改时间最新
    import time
    os.utime(os.path.join(month_dir, "19 3月AI筛选20日线汇总0305.xlsx"), (time.time(), time.time()))

    db_rows = [("2027-03-08", _compress([{"证券代码": "000009.SZ"}]))]
    with patch("core.database.AsyncSessionLocal", _fake_session_factory(db_rows)):
        result = await _run(_executor(base), "2027-03-08")
    assert result["success"], result.get("message")

    files = sorted(os.listdir(month_dir))
    assert files == [f"19 3月AI筛选20日线汇总{dd}.xlsx" for dd in ("0305", "0306", "0307", "0308")], files
    # 0308 = 复制0307 + 当日内容；无Sheet1旧文件 → 全量聚合构建 Sheet1
    wb = load_workbook(os.path.join(month_dir, "19 3月AI筛选20日线汇总0308.xlsx"))
    assert wb.sheetnames == ["Sheet1", "0308", "0307", "0306", "0305"], wb.sheetnames
    ws1 = wb["Sheet1"]
    assert [ws1.cell(1, c).value for c in range(1, 6)] == ["一级板块", "3月5日", "3月6日", "3月7日", "3月8日"]

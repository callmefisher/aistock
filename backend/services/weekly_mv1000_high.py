"""工作流：本周总市值千亿新高（14号表 + 20号表）。

规则（数据日期 date_str，回溯35天含当日）：
- 输出文件：{base_dir}/{date}/本周总市值千亿新高输出/14 总市值大于1000亿 本周最高价格大于上周最高价格{MMDD}.xlsx
  （注意与公共上传目录「本周总市值千亿新高」区分；同名强制覆盖更新）
- 目标文件解析：35天回溯（数据日期已有文件直接编辑；否则取最近35日内最新文件拷贝重命名；都没有则全新生成）；
- 全新生成：第一个 sheet「筛选结果」= 公共文件内容原样复制；
- 无论何种模式：新建以 MMDD 命名的第2个 sheet（位于「筛选结果」之后，重复执行先删后建保证幂等），
  内容为公共文件数据：按代码列（优先「股票代码」，其次「证券代码」，否则A列）降序整体排序，
  排序后在B列后插入一列空白列；
- 第一个 sheet「筛选结果」内容不修改；
- 最后将 MMDD sheet 的 A 列（含表头）复制到新 Excel（20号表），固定输出到
  {base_dir}/{date}/20号表本周总市值千亿新高A列/20 总市值大于1000亿 本周最高价格大于上周最高价格.xlsx。

复用 weekly_common：35天目标文件解析 / 公共数据源定位 / 行值读取。
"""

import logging
import re
import os
from typing import List, Optional

from openpyxl import Workbook, load_workbook

from services.weekly_common import (
    find_public_source_file,
    mmdd_of,
    read_sheet_rows,
    resolve_daily_output_file,
)
from utils.stock_code_normalizer import normalize_stock_code

logger = logging.getLogger(__name__)

# 14号表目标输出子目录（注意与公共上传目录「本周总市值千亿新高」区分）
OUT_SUBDIR = "本周总市值千亿新高输出"
# 公共数据源子目录
PUBLIC_SUBDIR = "本周总市值千亿新高"
# 默认输出文件名模板
DEFAULT_FILENAME_TEMPLATE = "14 总市值大于1000亿 本周最高价格大于上周最高价格{MMDD}.xlsx"
# 20号表输出目录名与固定文件名
TABLE20_SUBDIR = "20号表本周总市值千亿新高A列"
TABLE20_FILENAME = "20 总市值大于1000亿 本周最高价格大于上周最高价格.xlsx"

FIRST_SHEET_NAME = "筛选结果"


def build_default_filename(date_str: str) -> str:
    """数据日期 → 默认输出文件名，如 2026-10-15 → 14 总市值大于1000亿 本周最高价格大于上周最高价格1015.xlsx"""
    return DEFAULT_FILENAME_TEMPLATE.replace("{MMDD}", mmdd_of(date_str))


def resolve_output_filename(config: dict, date_str: str) -> str:
    """解析输出文件名：优先使用配置（前端可编辑），否则按数据日期生成默认名。"""
    configured = (config.get("output_filename") or "").strip()
    if configured:
        return configured.replace("{MMDD}", mmdd_of(date_str))
    return build_default_filename(date_str)


def _find_code_col_index(header: List) -> Optional[int]:
    """定位代码列索引：优先「股票代码」，其次「证券代码」，否则 None（由调用方回退A列）。"""
    for i, h in enumerate(header):
        if str(h).strip() == "股票代码":
            return i
    for i, h in enumerate(header):
        if str(h).strip() == "证券代码":
            return i
    return None


def sort_rows_by_code_desc(rows: List[List]) -> List[List]:
    """数据行按代码列降序整体排序（表头不动，其余列跟随不混行；无代码的行排在末尾）。"""
    if not rows:
        return []
    header = rows[0]
    data = [r for r in rows[1:] if any(v not in (None, "") for v in r)]
    code_idx = _find_code_col_index(header)

    def key(row: List):
        raw = row[code_idx] if code_idx is not None and code_idx < len(row) else (row[0] if row else None)
        nc = normalize_stock_code(raw)
        return (nc is not None, nc or "")

    data.sort(key=key, reverse=True)
    return [header] + data


def _is_invalid_data_row(row) -> bool:
    """数据行无效判定：全空行，或任一非空单元格为「数据来源」说明行 /
    undefined / null / none / nan 等脏数据（写入 MMDD sheet 前排除）。"""
    if not row or not any(v not in (None, "") for v in row):
        return True
    for v in row:
        if v in (None, ""):
            continue
        s = str(v).strip()
        if s.startswith("数据来源"):
            return True
        low = s.lower()
        if "undefined" in low or low in ("null", "none", "nan"):
            return True
    return False


def _write_mmdd_sheet(wb, sheet_name: str, rows: List[List]) -> None:
    """（重）建 MMDD sheet：位于「筛选结果」之后；内容=排序结果+B列后空白列。

    数据行排除「数据来源」/undefined 等无效行（表头保留原样；
    「筛选结果」sheet 不修改，保持公共文件原貌）。"""
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name, index=1)
    sorted_rows = sort_rows_by_code_desc(rows)
    filtered = [sorted_rows[0]] + [r for r in sorted_rows[1:] if not _is_invalid_data_row(r)]
    skipped = len(sorted_rows) - len(filtered)
    if skipped:
        logger.info(f"[本周总市值千亿新高] MMDD sheet 排除无效行 {skipped} 行（数据来源/undefined等）")
    for row in filtered:
        ws.append(["" if v is None else v for v in row])
    # B列后插入一列空白列
    ws.insert_cols(3)


def _is_invalid_table20_cell(v) -> bool:
    """20号表无效行判定：空值、「数据来源」开头的说明行、undefined/null 等脏数据。"""
    if v is None:
        return True
    s = str(v).strip()
    if not s:
        return True
    if s.startswith("数据来源"):
        return True
    low = s.lower()
    if "undefined" in low or low in ("null", "none", "nan"):
        return True
    return False


def _export_table20(base_dir: str, date_str: str, mv1000_path: str, mmdd: str) -> str:
    """20号表：从已生成的 14号表 的 MMDD sheet 拷贝第一列（A列，含表头）到独立 Excel，
    按股票代码/证券代码降序排列，仅此一列；排除「数据来源」/undefined 等无效行；固定文件名输出。"""
    out_dir = os.path.join(base_dir, date_str, TABLE20_SUBDIR)
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, TABLE20_FILENAME)

    # 读取 14号表 MMDD sheet 的 A 列（第一列）
    wb_src = load_workbook(mv1000_path, read_only=True, data_only=True)
    try:
        if mmdd not in wb_src.sheetnames:
            raise ValueError(f"14号表缺少 sheet「{mmdd}」: {mv1000_path}")
        ws_src = wb_src[mmdd]
        rows = [list(row) for row in ws_src.iter_rows(min_col=1, max_col=1, values_only=True)]
    finally:
        wb_src.close()

    if not rows:
        raise ValueError(f"14号表 sheet「{mmdd}」无数据: {mv1000_path}")

    # 表头保留原样；数据行剔除无效行（数据来源说明行、undefined/null 脏数据、空行）
    header, data_rows = rows[0], rows[1:]
    col_a = [header] + [r for r in data_rows if not _is_invalid_table20_cell(r[0] if r else None)]

    # 单列按代码降序（表头不动），复用既有排序逻辑
    sorted_col = sort_rows_by_code_desc(col_a)

    # 直接写入已知可导入 Choice 的标准 xlsx 包（单列，不经 openpyxl 序列化与事后修补）；
    # None 值归一为空串，由包构建函数写成带样式的空单元格
    _write_table20_package(
        out_path, mmdd,
        [row[0] if row and row[0] is not None else "" for row in sorted_col])
    return out_path


_TABLE20_TPL_CACHE: dict = {}


def _table20_template(filename: str) -> bytes:
    """20号表包模板部件（config/ 下，提取自 Choice 实测可导入的 WPS 模板），进程内缓存。

    filename: table20_styles_template.xml（样式表）/ table20_theme_template.xml（主题）
    """
    data = _TABLE20_TPL_CACHE.get(filename)
    if data is None:
        path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "config", filename)
        with open(path, "rb") as f:
            data = f.read()
        _TABLE20_TPL_CACHE[filename] = data
    return data


_XML_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'


def _write_table20_package(out_path: str, sheet_name: str, values: list) -> None:
    """直接构建 20号表 xlsx 包：单列字符串（含表头），不经 openpyxl 保存、无需事后修补。

    输出形态与「WPS 打开后原样重存」一致（Choice 等严格导入器实测可接受）：
    - 所有 xml/rels 部件带 <?xml?> 声明；[Content_Types].xml 为 zip 首条目（OPC 规范）
    - 关系 Target 全部相对路径；无空 workbookProtection/definedNames 元素
    - 字符串写 sharedStrings 共享字符串表（<c t="s">），不用 inlineStr 形态
    - styles/theme 用模板部件；样式链与模板一致：单元格 s="1" → cellXfs[1]
      numFmtId=176 fontId=0（宋体11、charset=134、主题色1、minor），A列整列
      style="1" 宽 18.75；默认行高 16.8、默认列宽 9（手工追加行自动继承列样式）
    """
    import zipfile

    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    # sharedStrings 去重建索引；空值写带样式的空单元格（不进共享表）
    sst_index: dict = {}
    sst_items: list = []
    rows_xml = []
    for i, v in enumerate(values, start=1):
        if v is None or v == "":
            rows_xml.append('<row r="%d" spans="1:1"><c r="A%d" s="1"/></row>' % (i, i))
            continue
        s = str(v)
        idx = sst_index.get(s)
        if idx is None:
            idx = len(sst_items)
            sst_index[s] = idx
            sst_items.append(s)
        rows_xml.append('<row r="%d" spans="1:1"><c r="A%d" s="1" t="s"><v>%d</v></c></row>'
                        % (i, i, idx))
    n_refs = sum(1 for v in values if v is not None and v != "")

    sheet = _XML_DECL + (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<dimension ref="A1:A%d"/>'
        '<sheetViews><sheetView tabSelected="1" workbookViewId="0">'
        '<selection activeCell="A1" sqref="A1"/></sheetView></sheetViews>'
        '<sheetFormatPr defaultColWidth="9" defaultRowHeight="16.8"/>'
        '<cols><col min="1" max="1" width="18.75" style="1" customWidth="1"/></cols>'
        '<sheetData>%s</sheetData>'
        '</worksheet>') % (len(values), "".join(rows_xml))
    sst = _XML_DECL + (
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        ' count="%d" uniqueCount="%d">%s</sst>'
        % (n_refs, len(sst_items),
           "".join('<si><t>%s</t></si>' % esc(t) for t in sst_items)))
    workbook = _XML_DECL + (
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<workbookPr/>'
        '<bookViews><workbookView activeTab="0"/></bookViews>'
        '<sheets><sheet name="%s" sheetId="1" r:id="rId1"/></sheets>'
        '<calcPr calcId="124519" fullCalcOnLoad="1"/>'
        '</workbook>') % esc(sheet_name)
    wb_rels = _XML_DECL + (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument'
        '/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument'
        '/2006/relationships/styles" Target="styles.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument'
        '/2006/relationships/theme" Target="theme/theme1.xml"/>'
        '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument'
        '/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
        '</Relationships>')
    root_rels = _XML_DECL + (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument'
        '/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006'
        '/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument'
        '/2006/relationships/extended-properties" Target="docProps/app.xml"/>'
        '</Relationships>')
    content_types = _XML_DECL + (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/xl/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
        '</Types>')
    app = _XML_DECL + (
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"'
        ' xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<Application>Microsoft Excel</Application>'
        '<HeadingPairs><vt:vector size="2" baseType="variant">'
        '<vt:variant><vt:lpstr>工作表</vt:lpstr></vt:variant>'
        '<vt:variant><vt:i4>1</vt:i4></vt:variant></vt:vector></HeadingPairs>'
        '<TitlesOfParts><vt:vector size="1" baseType="lpstr">'
        '<vt:lpstr>%s</vt:lpstr></vt:vector></TitlesOfParts>'
        '</Properties>') % esc(sheet_name)
    core = _XML_DECL + (
        '<cp:coreProperties'
        ' xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
        ' xmlns:dc="http://purl.org/dc/elements/1.1/">'
        '<dc:creator>aistock</dc:creator>'
        '</cp:coreProperties>')

    parts = [
        ("[Content_Types].xml", content_types.encode("utf-8")),
        ("_rels/.rels", root_rels.encode("utf-8")),
        ("docProps/app.xml", app.encode("utf-8")),
        ("docProps/core.xml", core.encode("utf-8")),
        ("xl/workbook.xml", workbook.encode("utf-8")),
        ("xl/_rels/workbook.xml.rels", wb_rels.encode("utf-8")),
        ("xl/styles.xml", _table20_template("table20_styles_template.xml")),
        ("xl/theme/theme1.xml", _table20_template("table20_theme_template.xml")),
        ("xl/sharedStrings.xml", sst.encode("utf-8")),
        ("xl/worksheets/sheet1.xml", sheet.encode("utf-8")),
    ]
    tmp = out_path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts:
            z.writestr(name, data)
    os.replace(tmp, out_path)


async def execute_weekly_mv1000_high(executor, config: dict, date_str: Optional[str] = None):
    """主入口：生成/更新 14号表「本周总市值千亿新高输出」+ 20号表 A列输出。"""
    date_str = (config.get("date_str") or "").strip() or (date_str or "").strip() or executor.today
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date_str):
        return {"success": False, "error": f"数据日期格式错误: {date_str}"}

    base_dir = executor.base_dir

    # 公共数据源（本周总市值千亿新高）
    public_path = find_public_source_file(base_dir, date_str, PUBLIC_SUBDIR)
    if not public_path:
        return {
            "success": False,
            "error": f"公共数据源缺失: {base_dir}/{date_str}/{PUBLIC_SUBDIR}/ 下无公共文件",
        }
    public_rows = read_sheet_rows(public_path)  # 优先「筛选结果」sheet，否则首个sheet
    if not public_rows:
        return {"success": False, "error": f"公共文件无数据: {public_path}"}

    output_filename = resolve_output_filename(config, date_str)
    target_path, mode = resolve_daily_output_file(
        base_dir, date_str, OUT_SUBDIR, output_filename, log_tag="本周总市值千亿新高"
    )
    logger.info(f"[本周总市值千亿新高] 目标文件: {target_path} (mode={mode})")

    mmdd = mmdd_of(date_str)

    if mode == "new":
        # 全新生成：第一个 sheet「筛选结果」= 公共文件内容原样复制
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = FIRST_SHEET_NAME
        for row in public_rows:
            ws.append(["" if v is None else v for v in row])
    else:
        # 编辑目标文件（reuse=同名可读 / copy=35日内拷贝重命名副本），「筛选结果」内容不修改。
        # 加载失败时降级为全新生成（目标文件解析已验证可读，此处兜底竞态/极端损坏，
        # 保证 20号表输出与步骤成功，不因单个坏文件中断工作流）。
        try:
            wb = load_workbook(target_path)
            if FIRST_SHEET_NAME not in wb.sheetnames:
                wb.worksheets[0].title = FIRST_SHEET_NAME
        except Exception as e:
            logger.warning(f"[本周总市值千亿新高] 目标文件加载失败，降级全新生成: {target_path}, {e}")
            wb = Workbook()
            ws = wb.active
            assert ws is not None
            ws.title = FIRST_SHEET_NAME
            for row in public_rows:
                ws.append(["" if v is None else v for v in row])
            mode = "new(降级)"

    # 只保留2个sheet：第1个「筛选结果」+ 第2个当前数据日期（如 1016）；
    # 拷贝来的历史文件可能残留旧日期sheet（1014/1015等），全部删除
    for name in list(wb.sheetnames):
        if name not in (FIRST_SHEET_NAME, mmdd):
            del wb[name]
    # 保证 sheet 顺序：筛选结果在前，数据日期在后
    if FIRST_SHEET_NAME in wb.sheetnames and wb.sheetnames[0] != FIRST_SHEET_NAME:
        wb.move_sheet(FIRST_SHEET_NAME, offset=-wb.sheetnames.index(FIRST_SHEET_NAME))

    _write_mmdd_sheet(wb, mmdd, public_rows)
    wb.save(target_path)
    wb.close()

    # 20号表：从已保存的 14号表 MMDD sheet 拷贝第一列（A列），降序，仅此一列
    table20_path = _export_table20(base_dir, date_str, target_path, mmdd)

    count = max(len(public_rows) - 1, 0)
    logger.info(f"[本周总市值千亿新高] 完成: {target_path} ({mode}) 记录数={count}; 20号表: {table20_path}")
    return {
        "success": True,
        "message": f"14号表已生成({mode})，数据日期 {date_str} 记录数 {count}；20号表A列已输出",
        "file_path": target_path,
        "table20_path": table20_path,
        # 20号表相对 {date} 的下载路径（含子目录），merge 进 step_config 供前端批量下载
        "_extra_outputs": {"output_filename_table20": f"{TABLE20_SUBDIR}/{TABLE20_FILENAME}"},
        "_preview": [{"数据日期": date_str, "记录数": count}],
    }

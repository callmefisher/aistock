"""本月AI筛选20日线汇总工作流实现（独立文件，复用「本月AI筛选汇总」共享逻辑）。

与「本月AI筛选汇总」(monthly_ai_summary.py) 的差异仅两处：
1. 数据源：「当月每日新增」输出文件的第2个sheet「20日线」（站上20日线当日新增）：
   - DB优先：workflow_results「当月每日新增」step_type='final_ma20'（save_workflow_result 追加写入）；
   - Excel兜底：data/excel/当月每日新增/{YYYY}年{M}月新增/*{MMDD}*.xlsx 的 sheet「20日线」；
2. 输出：data/excel/本月AI筛选20日线汇总/{YYYY-MM}/19 {M}月AI筛选20日线汇总{MMDD}.xlsx
   （目录为空新建/同名复用/无同名复制最新日期文件，规则同「本月AI筛选汇总」）。

其余逻辑完全复用 monthly_ai_summary 的实现（修 bug 一处生效）：
- 日期sheet（MMDD命名）：证券代码 + 一级板块 + 二级板块，按数据日期强制覆盖；
- Sheet1：一级/二级板块 × 日期 两张统计表 + 两张彩色折线图，增量更新（已有列/行保持原样，
  当日列刷新或追加，累计0留空），图表按新表尺寸重建；
- 全部A股板块匹配（共享缓存）；sheet重排（Sheet1, 日期降序, 其他）。

效率要点同 monthly_ai_summary：工作簿单次加载单次保存；向量化匹配；样式对象复用。
"""

import os
import re
import json
import zlib
import logging
from typing import Any, Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook

from utils.stock_code_normalizer import normalize_stock_code
# 共享逻辑复用（Sheet1增量更新/重排/目标文件解析/全部A股板块缓存），避免两份实现漂移
from services.monthly_ai_summary import (
    load_all_a_sector_map,
    _update_summary_sheet,
    _reorder_sheets,
    _resolve_output_path,
)

logger = logging.getLogger(__name__)

# 输出文件名：前缀19 + 关键字AI筛选20日线汇总（如 19 10月AI筛选20日线汇总1012.xlsx）
_OUTPUT_PREFIX = "19"
_OUTPUT_KEYWORD = "AI筛选20日线汇总"
# 输出目录名（data/excel/本月AI筛选20日线汇总/{YYYY-MM}/）
_OUTPUT_DIR_NAME = "本月AI筛选20日线汇总"


# ======================================================================
# 「当月每日新增」当日 20日线 数据读取（DB优先，Excel兜底）
# ======================================================================

async def _load_ma20_new_codes(base_dir: str, date_str: str) -> Optional[List[str]]:
    """读取数据日期当日的「当月每日新增」第2个sheet「20日线」证券代码。

    返回：None=完全无数据；[]=当日为空（合法，输出空sheet）；否则代码列表（去重保序）。
    """
    records = await _load_ma20_records_from_db(date_str)
    if records is None:
        records = _read_ma20_records_from_excel(base_dir, date_str)
    if records is None:
        return None

    codes: List[str] = []
    seen: Set[str] = set()
    for r in records:
        if not isinstance(r, dict):
            continue
        raw = str(r.get("证券代码") or "").strip()
        nc = normalize_stock_code(raw)
        if not nc or nc in seen:
            continue
        seen.add(nc)
        codes.append(raw)
    return codes


async def _load_ma20_records_from_db(date_str: str) -> Optional[List[dict]]:
    """DB读取：workflow_results「当月每日新增」step_type='final_ma20'（第2个sheet「20日线」）"""
    from core.database import AsyncSessionLocal
    from sqlalchemy import text

    try:
        async with AsyncSessionLocal() as session:
            result = await session.execute(
                text("""
                    SELECT data_compressed FROM workflow_results
                    WHERE workflow_type = '当月每日新增' AND step_type = 'final_ma20' AND date_str = :d
                    ORDER BY id DESC LIMIT 1
                """),
                {"d": date_str},
            )
            row = result.fetchone()
            if not row or not row[0]:
                return None
            try:
                records = json.loads(zlib.decompress(row[0]).decode("utf-8"))
            except Exception as e:
                logger.warning(f"[{_OUTPUT_KEYWORD}] DB记录解压失败({date_str}): {e}")
                return None
            return records if isinstance(records, list) else None
    except Exception as e:
        logger.warning(f"[{_OUTPUT_KEYWORD}] DB查询当月每日新增20日线失败({date_str}): {e}")
        return None


def _read_ma20_records_from_excel(base_dir: str, date_str: str) -> Optional[List[dict]]:
    """Excel兜底：data/excel/当月每日新增/{YYYY}年{M}月新增/ 下含 {MMDD} 的xlsx第2个sheet「20日线」"""
    try:
        mmdd = date_str[5:7] + date_str[8:10]
        month_dir = os.path.join(base_dir, "当月每日新增", f"{date_str[:4]}年{int(date_str[5:7])}月新增")
        if not os.path.isdir(month_dir):
            return None
        candidates = [f for f in os.listdir(month_dir) if f.endswith(".xlsx") and mmdd in f and not f.startswith("~")]
        if not candidates:
            return None
        path = max(candidates, key=lambda f: os.path.getmtime(os.path.join(month_dir, f)))
        df = pd.read_excel(os.path.join(month_dir, path), sheet_name="20日线", dtype=str)
        if "证券代码" not in df.columns:
            return None
        return df.to_dict("records")
    except ValueError:
        # 命中的文件无「20日线」sheet（当日20日线数据缺失，monthly_daily_new 跳过该sheet）
        logger.warning(f"[{_OUTPUT_KEYWORD}] Excel兜底文件无「20日线」sheet({date_str})")
        return None
    except Exception as e:
        logger.warning(f"[{_OUTPUT_KEYWORD}] Excel兜底读取失败({date_str}): {e}")
        return None


# ======================================================================
# 主步骤：本月AI筛选20日线汇总
# ======================================================================

async def execute_monthly_ai_ma20_summary(executor, config: Dict[str, Any], date_str: Optional[str] = None) -> Dict[str, Any]:
    """本月AI筛选20日线汇总主步骤（step_type: monthly_ai_ma20_summary）。

    executor: WorkflowExecutor 实例（使用其 base_dir / today）
    """
    date_str = (config.get("date_str") or "").strip() or date_str or executor.today
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date_str or ""):
        return {"success": False, "message": f"数据日期格式错误: {date_str}（需YYYY-MM-DD）"}

    mmdd = date_str[5:7] + date_str[8:10]
    month_int = int(date_str[5:7])
    month_root = date_str[:7]  # 如 2026-10

    # 1) 当日「当月每日新增」20日线数据（第2个sheet）
    codes = await _load_ma20_new_codes(executor.base_dir, date_str)
    if codes is None:
        return {
            "success": False,
            "message": f"数据日期({date_str})无「当月每日新增」20日线数据（DB与Excel均未找到），请先执行「当月每日新增」工作流",
        }

    # 2) 全部A股板块映射（与本月AI筛选汇总共享缓存）
    warnings: List[str] = []
    try:
        sector_map = load_all_a_sector_map(executor.base_dir)
    except Exception as e:
        sector_map = {}
        logger.warning(f"[{_OUTPUT_KEYWORD}] 读取全部A股失败: {e}")
    if not sector_map:
        warnings.append("未找到 data/excel/全部A股/ 有效文件，「一级板块/二级板块」列将为空")

    # 3) 日期sheet行：证券代码 + 一级板块 + 二级板块（向量化匹配）
    if sector_map and codes:
        norm_series = pd.Series([normalize_stock_code(c) for c in codes], dtype=object)
        l1_series = norm_series.map(lambda nc: sector_map.get(nc, ("", ""))[0])
        l2_series = norm_series.map(lambda nc: sector_map.get(nc, ("", ""))[1])
        rows = list(zip(codes, l1_series.tolist(), l2_series.tolist()))
        matched_l1 = int((l1_series != "").sum())
        matched_l2 = int((l2_series != "").sum())
    else:
        rows = [(c, "", "") for c in codes]
        matched_l1 = matched_l2 = 0

    # 4) 目标文件：data/excel/本月AI筛选20日线汇总/{YYYY-MM}/
    #    目录为空→新建；有同名→直接填充；非空无同名→复制最新日期文件为新名
    out_dir = os.path.join(executor.base_dir, _OUTPUT_DIR_NAME, month_root)
    os.makedirs(out_dir, exist_ok=True)
    default_name = f"{_OUTPUT_PREFIX} {month_int}月{_OUTPUT_KEYWORD}{mmdd}.xlsx"
    output_filename = (config.get("output_filename") or "").strip() or default_name
    output_filename = output_filename.replace("{date}", date_str).replace("{mmdd}", mmdd)
    if not output_filename.endswith(".xlsx"):
        output_filename += ".xlsx"
    target_path, path_mode = _resolve_output_path(
        out_dir, output_filename, prefix=_OUTPUT_PREFIX, keyword=_OUTPUT_KEYWORD, log_tag=_OUTPUT_KEYWORD
    )

    # 5) 打开/创建工作簿 → 强制覆盖日期sheet → 增量更新Sheet1 → 重排 → 保存
    file_existed = path_mode != "new"
    wb = load_workbook(target_path) if file_existed else Workbook()
    try:
        if not file_existed and wb.sheetnames == ["Sheet"]:
            wb.remove(wb.active)  # 移除默认空sheet，稍后由 Sheet1/日期sheet 重建

        # 日期sheet（强制覆盖同名）
        if mmdd in wb.sheetnames:
            del wb[mmdd]
        ws_date = wb.create_sheet(mmdd)
        ws_date.append(["证券代码", "一级板块", "二级板块"])
        for row in rows:
            ws_date.append(list(row))
        ws_date.column_dimensions["A"].width = 14
        ws_date.column_dimensions["B"].width = 18
        ws_date.column_dimensions["C"].width = 18
        ws_date.freeze_panes = "A2"

        # Sheet1 增量更新（已有列内容/行序保持，新板块追加，当日列追加或刷新，图表重建）
        _update_summary_sheet(wb, mmdd)

        # 重排：Sheet1, 日期sheet降序(1012, 1009, 1008...), 其他
        _reorder_sheets(wb)

        wb.save(target_path)
    except Exception as e:
        logger.error(f"[{_OUTPUT_KEYWORD}] 写入失败 {target_path}: {e}")
        return {"success": False, "message": f"生成Excel失败: {e}"}

    logger.info(f"[{_OUTPUT_KEYWORD}] 输出 {target_path} (sheet {mmdd}: {len(rows)}行, 模式={path_mode})")

    mode_label = {"new": "新建", "reuse": "复用已有", "copy": "基于最新文件复制为"}[path_mode]
    message = (
        f"本月AI筛选20日线汇总完成: 数据日期{date_str}，sheet「{mmdd}」{len(rows)}个证券代码，"
        f"一级板块匹配{matched_l1}个/二级板块匹配{matched_l2}个，"
        f"{mode_label}文件 {os.path.basename(target_path)}"
    )
    if warnings:
        message += "；" + "；".join(warnings)

    # 预览数据（日期sheet内容）
    from services.workflow_executor import clean_df_for_json  # 延迟导入避免循环依赖
    preview_df = pd.DataFrame(rows, columns=["证券代码", "一级板块", "二级板块"])
    return {
        "success": True,
        "message": message,
        "data": clean_df_for_json(preview_df),
        "columns": ["证券代码", "一级板块", "二级板块"],
        "rows": len(rows),
        "file_path": target_path,
        "warnings": warnings,
        "_df": preview_df,
    }

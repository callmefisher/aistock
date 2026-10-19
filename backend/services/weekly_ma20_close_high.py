"""工作流：本周20日线新高（13号表）。

规则（数据日期 date_str，回溯35天含当日）：
- 输出文件：{base_dir}/{date}/本周20日收盘价新高输出/13 占20日线无振幅{MMDD}.xlsx（同名强制覆盖更新）
- 数据源：{base_dir}/{date}/本周20日收盘价新高/ 下公共文件的数据行数量 = 股票数量
- 全新生成：第一个 sheet 生成表格（宋体12），表头「日期 / 股票数量」，日期列显示为「10月15日」样式
  （真实日期值 + m"月"d"日" 数字格式，点击可见完整年月日）；
- 已有文件：不修改已有数据行，仅按规则新增一行（数据日期已存在则刷新数量）；
- 表格更新完毕后（重）生成折线图：横轴为日期列（DateAxis 按周展示，majorUnit=7天，
  数据日期为最后一个数据点），纵轴为「股票数量」。

复用 weekly_common：35天目标文件解析 / 公共数据源定位。
"""

import logging
import re
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.marker import Marker
from openpyxl.styles import Font

from services.weekly_common import (
    find_public_source_file,
    mmdd_of,
    resolve_daily_output_file,
)

logger = logging.getLogger(__name__)

# 13号表目标输出子目录
OUT_SUBDIR = "本周20日收盘价新高输出"
# 公共数据源子目录
PUBLIC_SUBDIR = "本周20日收盘价新高"
# 默认输出文件名模板（MMDD 为占位符）
DEFAULT_FILENAME_TEMPLATE = "13 占20日线无振幅{MMDD}.xlsx"

TABLE_FONT = Font(name="宋体", size=12)
DATE_NUMBER_FORMAT = 'm"月"d"日"'
_SHEET_NAME = "Sheet1"


def build_default_filename(date_str: str) -> str:
    """数据日期 → 默认输出文件名，如 2026-10-15 → 13 占20日线无振幅1015.xlsx"""
    return DEFAULT_FILENAME_TEMPLATE.replace("{MMDD}", mmdd_of(date_str))


def resolve_output_filename(config: dict, date_str: str) -> str:
    """解析输出文件名：优先使用配置（前端可编辑），否则按数据日期生成默认名。"""
    configured = (config.get("output_filename") or "").strip()
    if configured:
        return configured.replace("{MMDD}", mmdd_of(date_str))
    return build_default_filename(date_str)


def count_public_rows(base_dir: str, date_str: str) -> Optional[int]:
    """统计公共文件（本周20日收盘价新高）的数据行数量。无公共文件时返回 None。"""
    import pandas as pd

    path = find_public_source_file(base_dir, date_str, PUBLIC_SUBDIR)
    if not path:
        return None
    df = pd.read_excel(path, dtype=str)
    return int(df.dropna(how="all").shape[0])


def _cell_to_date(value) -> Optional[date]:
    """单元格值 → date。

    兼容四种形态：
    - datetime / date 对象；
    - int/float Excel 日期序列值（手工/外部工具做的老文件常见：值如 46184、
      格式 General，openpyxl 读出为 int。若不识别会导致整列日期行判定为空，
      新行错误写入第2行覆盖历史数据、图表只剩单点）；
    - 『M月D日』文本；『YYYY-MM-DD』文本。

    序列值限定 2000~2100 年范围，防止误判普通数字。
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        from openpyxl.utils.datetime import from_excel
        try:
            dt = from_excel(value)
        except Exception:
            return None
        if isinstance(dt, datetime):
            dt = dt.date()
        if dt and 2000 <= dt.year <= 2100:
            return dt
        return None
    if isinstance(value, str):
        m = re.match(r"^\s*(\d{1,2})月(\d{1,2})日\s*$", value)
        if m:
            return date(datetime.now().year, int(m.group(1)), int(m.group(2)))
        m = re.match(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})", value)
        if m:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _table_font(ws, up_to_row: int = None) -> None:
    """表格区域统一字体：宋体 12（默认刷到最后一行；指定 up_to_row 时只刷到该行，不碰表格下方杂散内容）。"""
    max_row = up_to_row if up_to_row is not None else ws.max_row
    if max_row < 1:
        return
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=2):
        for cell in row:
            cell.font = TABLE_FONT


def _apply_table(ws, row_date: date, count: int) -> int:
    """写入/更新表格：表头 + 数据行（不修改已有行，数据日期已存在则仅刷新数量）。

    返回数据行数（含本次新增）。"""
    # 表头（全新文件 / 空表）
    if ws.max_row < 1 or ws["A1"].value in (None, ""):
        ws["A1"] = "日期"
        ws["B1"] = "股票数量"

    # 扫描日期列：定位「数据日期已有行」与「最后一个有效日期行」
    # （表格下方可能存在空行/杂散内容，追加必须紧贴最后一个日期行，而非 max_row）
    target_row = None
    last_date_row = None
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=1)
        cell_date = _cell_to_date(cell.value)
        if cell_date is not None:
            # 历史行可能为 int 序列值 + General 格式（手工文件），显示为数字 46184；
            # 统一刷为「X月X日」格式（不改值），保证表格展示一致
            if not isinstance(cell.value, (datetime, date)):
                cell.number_format = DATE_NUMBER_FORMAT
            last_date_row = r
            if cell_date == row_date:
                target_row = r

    if target_row is None:
        # 新行 = 表格里最后一个日期的下一行（同日已存在则只刷新数量）。
        # last_date_row=None 且表格非空时说明 A 列全是无法解析的杂散内容——
        # 此时不覆盖第2行（防丢数据），从 max_row+1 追加。
        if last_date_row is not None:
            target_row = last_date_row + 1
        else:
            target_row = ws.max_row + 1 if ws.max_row >= 1 else 2
        ws.cell(row=target_row, column=1).value = datetime(row_date.year, row_date.month, row_date.day)
        last_date_row = target_row

    ws.cell(row=target_row, column=2).value = count
    ws.cell(row=target_row, column=1).number_format = DATE_NUMBER_FORMAT

    _table_font(ws, last_date_row)
    return max((last_date_row or 1) - 1, 0)


def _rebuild_chart(ws, data_rows: int) -> None:
    """（重）生成折线图：横轴日期列（按周 majorUnit=7 天），纵轴股票数量。"""
    # 移除旧图，避免重复叠加
    ws._charts = []

    if data_rows <= 0:
        return

    chart = LineChart()
    chart.title = "股票数量趋势"
    chart.style = 2
    chart.height = 9
    chart.width = 24

    data = Reference(ws, min_col=2, min_row=1, max_row=1 + data_rows)
    cats = Reference(ws, min_col=1, min_row=2, max_row=1 + data_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    # 日期横轴：按周（7天）展示，避免时间变长后拥挤；数据日期为最后一个数据点
    chart.x_axis = DateAxis(crossAx=100)
    chart.x_axis.number_format = DATE_NUMBER_FORMAT
    chart.x_axis.majorUnit = 7
    chart.x_axis.majorTimeUnit = "days"
    chart.x_axis.delete = False
    chart.y_axis.crossAx = 500  # 与 DateAxis 默认 axId 对齐
    chart.y_axis.delete = False
    chart.y_axis.title = "股票数量"

    s = chart.series[0]
    s.smooth = False
    s.marker = Marker(symbol="circle", size=5)

    ws.add_chart(chart, "D2")


async def execute_weekly_ma20_close_high(executor, config: dict, date_str: Optional[str] = None):
    """主入口：生成/更新 13号表「本周20日收盘价新高输出」。"""
    date_str = (config.get("date_str") or "").strip() or (date_str or "").strip() or executor.today
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date_str):
        return {"success": False, "error": f"数据日期格式错误: {date_str}"}

    base_dir = executor.base_dir

    output_filename = resolve_output_filename(config, date_str)
    target_path, mode = resolve_daily_output_file(
        base_dir, date_str, OUT_SUBDIR, output_filename, log_tag="本周20日线新高"
    )
    logger.info(f"[本周20日线新高] 目标文件: {target_path} (mode={mode})")

    count = count_public_rows(base_dir, date_str)
    if count is None:
        return {
            "success": False,
            "error": f"公共数据源缺失: {base_dir}/{date_str}/{PUBLIC_SUBDIR}/ 下无公共文件",
        }

    row_date = datetime.strptime(date_str, "%Y-%m-%d").date()

    if mode == "new":
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = _SHEET_NAME
    else:
        # reuse/copy 模式：目标文件解析已验证可读，此处 try 兜底竞态/极端损坏，降级全新生成
        try:
            wb = load_workbook(target_path)
            ws = wb.worksheets[0]
        except Exception as e:
            logger.warning(f"[本周20日线新高] 目标文件加载失败，降级全新生成: {target_path}, {e}")
            wb = Workbook()
            ws = wb.active
            assert ws is not None
            ws.title = _SHEET_NAME
            mode = "new(降级)"

    data_rows = _apply_table(ws, row_date, count)
    _rebuild_chart(ws, data_rows)
    wb.save(target_path)
    wb.close()

    logger.info(f"[本周20日线新高] 完成: {target_path} 股票数量={count} 数据行数={data_rows}")
    return {
        "success": True,
        "message": f"13号表已生成({mode})，数据日期 {date_str} 股票数量 {count}",
        "file_path": target_path,
        "_preview": [{"数据日期": date_str, "股票数量": count}],
    }
